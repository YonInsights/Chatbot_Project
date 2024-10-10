import streamlit as st
import fitz  # PyMuPDF
import docx
import os
import sys
from fpdf import FPDF

# Add the project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from chatbot.chatbot_logic import start_chatbot

def main():
    st.title("Skill Roadmap Chatbot")
    
    # CV Upload
    uploaded_file = st.file_uploader("Upload your CV", type=["pdf", "docx"])
    
    # User Inputs
    st.write("Please answer the following questions:")
    goal = st.text_input("What new skill or field are you interested in learning?")
    time_available = st.number_input("How much time per week can you dedicate to learning this new skill?", min_value=1, max_value=40, step=1)
    
    if st.button("Generate Roadmap"):
        if uploaded_file is not None:
            # Extract text from the uploaded file
            cv_content = extract_text_from_file(uploaded_file)
            
            # Call the chatbot function with the extracted CV content and user inputs
            roadmap = start_chatbot(cv_content, goal, time_available)
            
            # Display the roadmap in a table
            st.write("Here is your personalized roadmap:")
            st.table({
                "Skill": roadmap["skills_to_learn"],
                "Suggested Material": roadmap["learning_materials"],
                "Estimated Time": roadmap["time_estimations"]
            })
            
            # Add motivational section
            st.write("Remember, consistency is key to mastering new skills. Believe in yourself and stay dedicated. You can do this!")
            
            # Generate PDF
            if st.button("Download as PDF"):
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                
                pdf.cell(200, 10, txt="Personalized Study Roadmap", ln=True, align="C")
                pdf.ln(10)
                
                for skill, material, time in zip(roadmap["skills_to_learn"], roadmap["learning_materials"], roadmap["time_estimations"]):
                    pdf.cell(200, 10, txt=f"Skill: {skill}", ln=True)
                    pdf.cell(200, 10, txt=f"Suggested Material: {material}", ln=True)
                    pdf.cell(200, 10, txt=f"Estimated Time: {time}", ln=True)
                    pdf.ln(5)
                
                pdf.ln(10)
                pdf.cell(200, 10, txt="Motivational Section:", ln=True)
                pdf.multi_cell(0, 10, txt="Remember, consistency is key to mastering new skills. Believe in yourself and stay dedicated. You can do this!")
                
                pdf_output_path = "roadmap.pdf"
                pdf.output(pdf_output_path)
                st.write(f"PDF generated! [Download it here](roadmap.pdf)")
        else:
            st.write("Please upload your CV.")

def extract_text_from_file(uploaded_file):
    if uploaded_file.type == "application/pdf":
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx(uploaded_file)
    else:
        return "Unsupported file type."

def extract_text_from_pdf(file):
    document = fitz.open(stream=file.read(), filetype="pdf")
    text = ""
    for page in document:
        text += page.get_text()
    return text

def extract_text_from_docx(file):
    doc = docx.Document(file)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

if __name__ == "__main__":
    main()
